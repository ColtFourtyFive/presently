from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'research'
OUT = ROOT / 'output' / 'documents'
OUT.mkdir(parents=True, exist_ok=True)

def diagram():
    im = Image.new('RGB', (1500, 810), 'white')
    d = ImageDraw.Draw(im)
    fontpath = '/System/Library/Fonts/Supplemental/Arial.ttf'
    boldpath = '/System/Library/Fonts/Supplemental/Arial Bold.ttf'
    title = ImageFont.truetype(boldpath, 29)
    body = ImageFont.truetype(fontpath, 26)
    small = ImageFont.truetype(fontpath, 23)
    ink = '#172A3A'
    def box(x,y,w,h,head,lines,fill='#F1F5F8'):
        d.rounded_rectangle((x,y,x+w,y+h),radius=10,fill=fill,outline='#8396A4',width=2)
        d.text((x+24,y+20),head,font=title,fill=ink)
        for j,line in enumerate(lines):
            d.text((x+24,y+62+j*34),line,font=body,fill=ink)
    def arrow(points):
        d.line(points,fill='#526F83',width=4)
        (x1,y1),(x2,y2)=points[-2:]
        import math
        a=math.atan2(y2-y1,x2-x1)
        d.polygon([(x2,y2),(x2-15*math.cos(a-.45),y2-15*math.sin(a-.45)),(x2-15*math.cos(a+.45),y2-15*math.sin(a+.45))],fill='#526F83')
    box(35,30,440,160,'Staff web application',['CRM and online attendance','Named staff and managed devices'])
    box(515,30,380,160,'Managed identity',['Staff authentication','MFA and session validation'])
    box(950,30,515,190,'Contingency companion',['One designated managed device','Encrypted roster and event queue','Bounded offline staff access'])
    box(240,310,715,155,'Application service',['Permissions and business rules','Transactional attendance and CRM commands'])
    arrow([(255,190),(255,260),(390,260),(390,310)])
    arrow([(705,190),(705,310)])
    arrow([(1110,220),(1110,385),(955,385)])
    d.text((1130,276),'Snapshot and',font=small,fill=ink)
    d.text((1130,307),'queued event sync',font=small,fill=ink)
    box(35,575,440,160,'PostgreSQL',['Student and attendance records','Audit and transactional outbox'])
    box(515,575,440,160,'Background worker',['Reports and record lifecycle','Later provider integrations'])
    box(995,575,470,160,'Encrypted object storage',['Scoped exports and attachments','Separate access and expiry'])
    arrow([(420,465),(420,515),(255,515),(255,575)])
    arrow([(735,465),(735,575)])
    arrow([(955,655),(995,655)])
    d.text((45,765),'Managed backups and monitoring support the primary service. Staff procedures govern contingency use.',font=small,fill=ink)
    im.save(SRC/'architecture.png')

def source_register():
    return '''SRC01. User-supplied Student Check-In and Check-Out System Requirements and Non-Exhaustive Informational Vendor List. Undated, two-page PDF, reviewed 14 September 2026. Page 1 contains eight baseline attendance requirements and the conditional future certification statement. Page 2 lists ten illustrative vendors, non-endorsement language, and illustrative pricing. The original attachment remains the reference source.

V01 to V10. Official vendor product, help, pricing and documentation pages listed with each vendor in the research appendix. All were accessed on 14 September 2026. Suffixes such as V01a identify an exact page. The research describes published claims and explicitly identifies unknowns.

REG01. US Federal Trade Commission. Complying with COPPA frequently asked questions. Applicability discussion reviewed 14 September 2026. https://www.ftc.gov/business-guidance/resources/complying-coppa-frequently-asked-questions

REG02. US Department of Education. To which educational agencies or institutions does FERPA apply. Applicability discussion reviewed 14 September 2026. https://studentprivacy.ed.gov/faq/which-educational-agencies-or-institutions-does-ferpa-apply

No corporate Kumon policy portal, franchise agreement, internal student-system documentation, vendor contract, or existing center dataset was supplied. Current policy status, legal applicability, plan-specific capabilities and integration permission remain decisions to confirm.'''

def appendix():
    chunks=[]
    for name in ['vendor_appendix_a.md','vendor_appendix_b.md']:
        s=(SRC/name).read_text()
        idx=s.index('### V')
        s=s[idx:]
        # Make readable link labels while preserving exact URL targets.
        lines=[]
        for line in s.splitlines():
            if line.startswith('- V'):
                m=re.match(r'- (V\d+[a-z])[:.]\s*(.*?)\s*(https?://\S+)\s*$',line)
                if m:
                    sid,label,url=m.groups()
                    labels={'V01a':'Attendance tracking','V01b':'Frequently asked questions','V02a':'Attendance software','V02b':'Pricing','V03a':'Child care management features','V03b':'Security controls','V03c':'Request pricing','V04a':'Attendance tracking','V04b':'Security FAQs','V04c':'Pricing','V05a':'Attendance and reporting','V05b':'Roster check-in','V05c':'Admin console','V05d':'Pricing'}
                    label=label.rstrip(': ').strip() or labels.get(sid,'Official vendor page')
                    line=f'- {sid}. [{label}]({url})'
            lines.append(line)
        chunks.append('\n'.join(lines).strip())
    return '\n\n'.join(chunks)

def hyperlink(p,label,url):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    col=OxmlElement('w:color');col.set(qn('w:val'),'1F4E79');props.append(col)
    ul=OxmlElement('w:u');ul.set(qn('w:val'),'single');props.append(ul)
    run.append(props);t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);p._p.append(link)

def inline(p,s):
    rx=r'\[([^\]]+)\]\((https?://[^)]+)\)|(https?://\S+)|\*\*([^*]+)\*\*|`([^`]+)`'
    pos=0
    for m in re.finditer(rx,s):
        if m.start()>pos:p.add_run(s[pos:m.start()])
        if m.group(1):hyperlink(p,m.group(1),m.group(2))
        elif m.group(3):hyperlink(p,m.group(3),m.group(3))
        elif m.group(4):p.add_run(m.group(4)).bold=True
        else:
            r=p.add_run(m.group(5));r.font.name='Courier New';r.font.size=Pt(10)
        pos=m.end()
    if pos<len(s):p.add_run(s[pos:])

def borders(table):
    pr=table._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
    pr.append(b)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',90),('bottom',90),('left',105),('right',105)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)

def table_widths(header,rows):
    n=len(header);total=6.9
    h=' '.join(header)
    if n==2:return [.62,6.28] if header[0]=='Test' else [2.05,4.85]
    if n==3:
        if header[0]=='Endpoint':return [2.2,2.15,2.55]
        if header[0] in ['ID','Test']:return [0.60,3.15,3.15]
        if header[0]=='Source item':return [.65,3.1,3.15]
        return [1.4,2.75,2.75]
    if n==4:
        if header[0]=='ID':return [.55,3.25,1.45,1.65] if header[1]=='Decision' else [.55,2.8,1.25,2.3]
        if header[0]=='Resource or action':return [2.45,1.05,1.1,2.3]
        return [1.0,2.55,1.55,1.8]
    if n==5:
        if header[0]=='Requirement':return [1.0,.90,2.05,1.35,1.6]
        return [2.22,.80,.91,1.49,1.48]
    return [total/n]*n

def add_table(doc,lines):
    rows=[[c.strip() for c in ln.strip().strip('|').split('|')] for ln in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in r)]
    header=rows[0]; widths=table_widths(header,rows[1:])
    t=doc.add_table(rows=0,cols=len(header));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders(t)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'244761' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'));pr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.06
            p.paragraph_format.keep_with_next=False
            if len(txt)<20 and j==0 and header[0] in ['ID','Source item','Requirement','Test']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,txt)
            for r in p.runs:
                r.font.size=Pt(10.2);r.font.name='Arial';r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000');r.bold=i==0
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=0.3;p.add_run().font.size=Pt(3)

def build(kind):
    raw=(SRC/f'{kind}.md').read_text()
    if kind=='BRD':raw=raw.replace('{{VENDOR_APPENDIX}}',appendix()).replace('{{SOURCE_REGISTER}}',source_register())
    raw=raw.translate(str.maketrans({'“':'"','”':'"','‘':"'",'’':"'"}))
    if '{{' in raw:raise ValueError('Unresolved placeholder')
    (OUT/f'Kumon_CRM_{kind}.md').write_text(raw.replace('(architecture.png)',f'({SRC / "architecture.png"})'))
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    sec.header_distance=Inches(.30);sec.footer_distance=Inches(.30)
    styles=d.styles
    for border in list(styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for border in list(d.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for nm in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
        styles[nm].font.name='Arial';styles[nm].font.color.rgb=RGBColor(0,0,0)
    normal=styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.10;normal.paragraph_format.space_after=Pt(7)
    for nm,size,before,after in [('Title',24,0,9),('Subtitle',11,0,8),('Heading 1',16,17,8),('Heading 2',12.5,12,6),('Heading 3',11.5,10,5)]:
        st=styles[nm];st.font.size=Pt(size);st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.keep_with_next=True
        if nm.startswith('Heading'):st.font.bold=True
    for st in styles:
        if st.type==1:
            try:st.paragraph_format.widow_control=True
            except Exception:pass
    h=sec.header.paragraphs[0];h.text=f'KUMON CENTER CRM  |  {kind}';h.style=styles['Normal'];h.paragraph_format.space_after=Pt(0)
    for r in h.runs:r.font.size=Pt(8.5);r.font.color.rgb=RGBColor(0,0,0)
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('Draft 0.1  |  14 September 2026  |  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    for r in f.runs:r.font.size=Pt(8.5)
    d.core_properties.title=raw.splitlines()[0][2:]
    d.core_properties.subject=f'Kumon center in-house CRM {kind}'
    d.core_properties.author='Prepared for the Kumon center owner'
    d.core_properties.keywords='Kumon, CRM, attendance, requirements, draft'
    lines=raw.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            batch=[]
            while i<len(lines) and lines[i].strip().startswith('|'):batch.append(lines[i]);i+=1
            add_table(d,batch);continue
        if line.startswith('# '):d.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):d.add_paragraph(line[3:],'Heading 1')
        elif line.startswith('### '):d.add_paragraph(line[4:],'Heading 2')
        elif line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line);p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(SRC/m.group(2)),width=Inches(6.9))
        elif line.startswith('- '):
            p=d.add_paragraph(style='List Bullet');inline(p,line[2:]);p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.keep_with_next=i+1<len(lines) and lines[i+1].strip().startswith('- ')
        elif line.startswith('Version 0.1'):
            d.add_paragraph(line,'Subtitle')
        else:
            p=d.add_paragraph();inline(p,line)
        i+=1
    dest=OUT/f'Kumon_CRM_{kind}.docx';d.save(dest)
    print(dest)

if __name__=='__main__':
    diagram()
    for k in ['BRD','FRD','FDR']:build(k)
